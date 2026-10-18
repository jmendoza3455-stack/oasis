"""Genera docs/YYMMDD Modelo economico Oasis.xlsx (OASIS.md §11.9).

Uso: python3 scripts/modelo-economico.py [YYMMDD]

Todas las celdas de supuestos están en la hoja Parámetros (amarillas) y el resto del libro son
fórmulas vivas que las referencian. Las cifras son supuestos marcados en docs/supuestos.md: nada
de esto es un dato real del negocio hasta que André/Carlos lo validen.
"""

import sys
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
STAMP = sys.argv[1] if len(sys.argv) > 1 else date.today().strftime("%y%m%d")
OUT = ROOT / "docs" / f"{STAMP} Modelo economico Oasis.xlsx"

GREEN = "2E3B2B"
SAND = "F7F5F0"
YELLOW = PatternFill("solid", fgColor="FFF2CC")
HEAD = PatternFill("solid", fgColor=GREEN)
SUB = PatternFill("solid", fgColor="E8E4D8")
THIN = Side(style="thin", color="C9C2B0")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
PEN = '"S/" #,##0;[Red]-"S/" #,##0'
PEN2 = '"S/" #,##0.00;[Red]-"S/" #,##0.00'
PCT = "0.0%"
NUM = "#,##0.0"

wb = Workbook()


def header(ws, row, values, widths=None):
    for i, v in enumerate(values, start=1):
        c = ws.cell(row=row, column=i, value=v)
        c.fill = HEAD
        c.font = Font(bold=True, color=SAND)
        c.alignment = Alignment(vertical="center", wrap_text=True)
        c.border = BOX
    if widths:
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w


def title(ws, text, note=None):
    ws["A1"] = text
    ws["A1"].font = Font(bold=True, size=14, color=GREEN)
    if note:
        ws["A2"] = note
        ws["A2"].font = Font(italic=True, size=9, color="6B6B6B")


# ---------------------------------------------------------------------------
# Léeme
# ---------------------------------------------------------------------------
readme = wb.active
readme.title = "Léeme"
title(readme, "Modelo económico · Oasis Pachacámac · Un espacio SAWA")
lines = [
    f"Versión {STAMP}. Generado con scripts/modelo-economico.py (no editar la estructura a mano: regenerar).",
    "",
    "Cómo usarlo",
    "1. Cambia SOLO las celdas amarillas de la hoja Parámetros. Todo lo demás son fórmulas vivas.",
    "2. Unit economics: margen por evento según segmento, tipo de día y canal.",
    "3. Mensual: 12 meses desde la reapertura con la rampa de ocupación de las metas (S10).",
    "4. Equilibrio: eventos y ocupación de equilibrio, y payback del CAPEX.",
    "5. Sensibilidad: resultado mensual con ocupación de 20% a 80% contra ticket ±20%.",
    "6. Pauta vs planner: cuánto puede costar como máximo un cliente directo frente al 5% del planner.",
    "",
    "Advertencias",
    "• Todas las cifras son SUPUESTOS (ver docs/supuestos.md). Tarifas = tarifario de ejemplo (S24), por validar.",
    "• Costos fijos y CAPEX marcados [COMPLETAR] deben reemplazarse por cifras reales.",
    "• Montos SIN IGV salvo que se indique. Impuesto a la renta y tratamiento tributario: por validar con contador.",
    "• No es una proyección financiera auditada; es una herramienta para decidir precios, metas y canales.",
]
for i, line in enumerate(lines, start=3):
    readme.cell(row=i, column=1, value=line)
    if line in ("Cómo usarlo", "Advertencias"):
        readme.cell(row=i, column=1).font = Font(bold=True, color=GREEN)
readme.column_dimensions["A"].width = 110

# ---------------------------------------------------------------------------
# Parámetros
# ---------------------------------------------------------------------------
par = wb.create_sheet("Parámetros")
title(par, "Parámetros (celdas amarillas = supuestos editables)", "Montos en S/ sin IGV salvo indicación.")
header(par, 4, ["Grupo", "Parámetro", "Valor", "Unidad", "Fuente / supuesto"], [22, 46, 14, 16, 70])

P: dict[str, str] = {}
row = 5


def p(key, group, label, value, unit, note, fmt=None):
    global row
    par.cell(row=row, column=1, value=group)
    par.cell(row=row, column=2, value=label)
    c = par.cell(row=row, column=3, value=value)
    c.fill = YELLOW
    c.border = BOX
    if fmt:
        c.number_format = fmt
    par.cell(row=row, column=4, value=unit)
    par.cell(row=row, column=5, value=note)
    P[key] = f"Parámetros!$C${row}"
    row += 1


def gap():
    global row
    row += 1


p("igv", "General", "IGV", 0.18, "%", "settings.igv_rate", PCT)
p("usd", "General", "Tipo de cambio referencial", 3.75, "S/ por USD", "settings.currency_usd_rate", "0.00")
p("ir", "General", "Impuesto a la renta (régimen general)", 0.295, "%", "Por validar con contador", PCT)
gap()
p("d_vsd", "Calendario", "Días V-S-D + feriados por mes", 13.8, "días/mes", "(52×3 + ~10 feriados) / 12", NUM)
p("d_lj", "Calendario", "Días L-J por mes", 16.6, "días/mes", "(365 − 166) / 12", NUM)
p("d_block", "Calendario", "Días de mantenimiento por mes", 1, "días/mes", "Bloqueos de mantenimiento (se descuentan de L-J)", NUM)
gap()
p("o_vsd3", "Metas", "Ocupación V-S-D meses 1–3", 0.5, "%", "S10", PCT)
p("o_vsd6", "Metas", "Ocupación V-S-D desde el mes 4", 0.7, "%", "S10", PCT)
p("o_lj6", "Metas", "Ocupación L-J al mes 6 (rampa lineal)", 0.15, "%", "S10", PCT)
p("high", "Metas", "Factor de ticket en temporada alta (dic–mar)", 1.12, "×", "Diferencia alta/baja del tarifario de ejemplo", "0.00")
gap()
SEGS = [("boda", "Boda / social"), ("corp", "Corporativo"), ("fd", "Full day familiar")]
for k, label, v in [("boda", "Boda / social", 0.5), ("corp", "Corporativo", 0.15), ("fd", "Full day familiar", 0.35)]:
    p(f"mix_vsd_{k}", "Mezcla V-S-D", label, v, "% de eventos", "Supuesto; debe sumar 100%", PCT)
for k, label, v in [("boda", "Boda / social", 0.1), ("corp", "Corporativo", 0.6), ("fd", "Full day familiar", 0.3)]:
    p(f"mix_lj_{k}", "Mezcla L-J", label, v, "% de eventos", "Supuesto; debe sumar 100%", PCT)
gap()
for k, label, v in [("boda", "Boda / social", 4233), ("corp", "Corporativo", 3933), ("fd", "Full day familiar", 2167)]:
    p(f"t_vsd_{k}", "Ticket V-S-D", label, v, "S/ sin IGV", "Promedio V-S-D del tarifario de ejemplo (S24)", PEN)
for k, label, v in [("boda", "Boda / social", 2700), ("corp", "Corporativo", 3000), ("fd", "Full day familiar", 1350)]:
    p(f"t_lj_{k}", "Ticket L-J", label, v, "S/ sin IGV", "Promedio L-J del tarifario de ejemplo (S24)", PEN)
gap()
p("ch_pl", "Canales", "Eventos vía planner", 0.4, "% de eventos", "Supuesto", PCT)
p("ch_dir", "Canales", "Eventos directos por pauta", 0.35, "% de eventos", "Supuesto", PCT)
p("ch_org", "Canales", "Eventos referidos / orgánicos", 0.25, "% de eventos", "Supuesto; debe sumar 100%", PCT)
p("com", "Canales", "Comisión del planner", 0.05, "% alquiler sin IGV", "S04 / settings.planner_commission_rate", PCT)
p("ads", "Canales", "Pauta mensual", 1500, "S/ por mes", "S09 / settings.monthly_ad_budget", PEN)
p("conv", "Canales", "Conversión lead directo → evento", 0.08, "%", "Supuesto; medir en el dashboard", PCT)
gap()
p("dc", "Costo directo", "Costo directo por evento", 1210, "S/ por evento", "S28: limpieza, personal, seguridad, energía, agua, piscina, reposiciones", PEN)
gap()
OPEX = [
    ("op_staff", "Encargado en sitio (guardián)", 1800),
    ("op_pool", "Mantenimiento de piscina", 500),
    ("op_garden", "Jardinería", 600),
    ("op_util", "Servicios básicos fijos", 400),
    ("op_ins", "Seguro de responsabilidad civil", 250),
    ("op_acc", "Contabilidad", 600),
    ("op_soft", "Software, web y dominio", 150),
    ("op_tel", "Internet y telefonía", 120),
    ("op_tax", "Arbitrios y predial", 250),
    ("op_other", "Otros e imprevistos", 300),
]
for k, label, v in OPEX:
    p(k, "OPEX fijo mensual", label, v, "S/ por mes", "[COMPLETAR] supuesto por validar", PEN)
gap()
CAPEX = [
    ("cx_rem", "Remodelación", 80000),
    ("cx_furn", "Mobiliario y menaje", 15000),
    ("cx_pool", "Equipamiento de piscina y seguridad", 8000),
    ("cx_brand", "Señalética y marca", 3000),
    ("cx_launch", "Lanzamiento", 5000),
]
for k, label, v in CAPEX:
    p(k, "CAPEX (una vez)", label, v, "S/", "[COMPLETAR] supuesto por validar", PEN)

# Controles
row += 1
par.cell(row=row, column=2, value="Control: mezcla V-S-D suma 100%").font = Font(italic=True)
par.cell(row=row, column=3, value=f"=IF(ABS({P['mix_vsd_boda']}+{P['mix_vsd_corp']}+{P['mix_vsd_fd']}-1)<0.001,\"OK\",\"REVISAR\")")
row += 1
par.cell(row=row, column=2, value="Control: mezcla L-J suma 100%").font = Font(italic=True)
par.cell(row=row, column=3, value=f"=IF(ABS({P['mix_lj_boda']}+{P['mix_lj_corp']}+{P['mix_lj_fd']}-1)<0.001,\"OK\",\"REVISAR\")")
row += 1
par.cell(row=row, column=2, value="Control: canales suman 100%").font = Font(italic=True)
par.cell(row=row, column=3, value=f"=IF(ABS({P['ch_pl']}+{P['ch_dir']}+{P['ch_org']}-1)<0.001,\"OK\",\"REVISAR\")")
par.freeze_panes = "A5"

# Derivados (fórmulas) reutilizables
OPEX_SUM = "+".join(P[k] for k, _, _ in OPEX)
CAPEX_SUM = "+".join(P[k] for k, _, _ in CAPEX)
T_VSD = "+".join(f"{P[f'mix_vsd_{k}']}*{P[f't_vsd_{k}']}" for k, _ in SEGS)
T_LJ = "+".join(f"{P[f'mix_lj_{k}']}*{P[f't_lj_{k}']}" for k, _ in SEGS)

# ---------------------------------------------------------------------------
# Unit economics
# ---------------------------------------------------------------------------
ue = wb.create_sheet("Unit economics")
title(ue, "Unit economics por evento", "Margen de contribución = alquiler sin IGV − costo directo − costo del canal.")
header(ue, 4, ["Segmento", "Día", "Ticket sin IGV", "IGV", "Ticket con IGV", "Costo directo", "Contribución antes de canal", "Margen %"], [22, 10, 15, 13, 15, 14, 18, 11])
r = 5
for k, label in SEGS:
    for day, pref in (("V-S-D", "t_vsd"), ("L-J", "t_lj")):
        ue.cell(row=r, column=1, value=label)
        ue.cell(row=r, column=2, value=day)
        ue.cell(row=r, column=3, value=f"={P[f'{pref}_{k}']}").number_format = PEN
        ue.cell(row=r, column=4, value=f"=C{r}*{P['igv']}").number_format = PEN
        ue.cell(row=r, column=5, value=f"=C{r}+D{r}").number_format = PEN
        ue.cell(row=r, column=6, value=f"={P['dc']}").number_format = PEN
        ue.cell(row=r, column=7, value=f"=C{r}-F{r}").number_format = PEN
        ue.cell(row=r, column=8, value=f"=IF(C{r}>0,G{r}/C{r},0)").number_format = PCT
        r += 1

r += 1
ue.cell(row=r, column=1, value="Estado estable (desde el mes 6)").font = Font(bold=True, color=GREEN)
r += 1
base = r
rows_ss = [
    ("Eventos V-S-D por mes", f"={P['d_vsd']}*{P['o_vsd6']}", NUM),
    ("Eventos L-J por mes", f"=({P['d_lj']}-{P['d_block']})*{P['o_lj6']}", NUM),
    ("Eventos totales por mes", f"=C{base}+C{base + 1}", NUM),
    ("Ticket promedio V-S-D", f"={T_VSD}", PEN),
    ("Ticket promedio L-J", f"={T_LJ}", PEN),
    ("Ticket promedio ponderado", f"=IF(C{base + 2}>0,(C{base}*C{base + 3}+C{base + 1}*C{base + 4})/C{base + 2},0)", PEN),
    ("Eventos directos por mes", f"=C{base + 2}*{P['ch_dir']}", NUM),
    ("CAC directo (pauta / eventos directos)", f"=IF(C{base + 6}>0,{P['ads']}/C{base + 6},0)", PEN),
]
for label, formula, fmt in rows_ss:
    ue.cell(row=r, column=1, value=label)
    c = ue.cell(row=r, column=3, value=formula)
    c.number_format = fmt
    r += 1
SS = {
    "e_vsd": f"'Unit economics'!$C${base}",
    "e_lj": f"'Unit economics'!$C${base + 1}",
    "e_tot": f"'Unit economics'!$C${base + 2}",
    "t_vsd": f"'Unit economics'!$C${base + 3}",
    "t_lj": f"'Unit economics'!$C${base + 4}",
    "ticket": f"'Unit economics'!$C${base + 5}",
    "cac": f"'Unit economics'!$C${base + 7}",
}

r += 1
header_row = r
for i, v in enumerate(["Canal", "", "Ticket promedio", "Costo del canal", "Costo directo", "Contribución", "Margen %", "% de eventos"], start=1):
    c = ue.cell(row=r, column=i, value=v)
    c.fill = SUB
    c.font = Font(bold=True)
r += 1
channel_rows = []
for label, cost, share in (
    ("Planner (5% comisión)", f"=C{{r}}*{P['com']}", P["ch_pl"]),
    ("Directo (pauta, CAC)", f"={SS['cac']}", P["ch_dir"]),
    ("Referido / orgánico", "=0", P["ch_org"]),
):
    ue.cell(row=r, column=1, value=label)
    ue.cell(row=r, column=3, value=f"={SS['ticket']}").number_format = PEN
    ue.cell(row=r, column=4, value=cost.format(r=r)).number_format = PEN
    ue.cell(row=r, column=5, value=f"={P['dc']}").number_format = PEN
    ue.cell(row=r, column=6, value=f"=C{r}-D{r}-E{r}").number_format = PEN
    ue.cell(row=r, column=7, value=f"=IF(C{r}>0,F{r}/C{r},0)").number_format = PCT
    ue.cell(row=r, column=8, value=f"={share}").number_format = PCT
    channel_rows.append(r)
    r += 1
ue.cell(row=r, column=1, value="Promedio ponderado por canal").font = Font(bold=True)
first, last = channel_rows[0], channel_rows[-1]
ue.cell(row=r, column=6, value=f"=SUMPRODUCT(F{first}:F{last},H{first}:H{last})").number_format = PEN
SS["contrib"] = f"'Unit economics'!$F${r}"
ue.freeze_panes = "A5"

# ---------------------------------------------------------------------------
# Mensual
# ---------------------------------------------------------------------------
mo = wb.create_sheet("Mensual")
title(mo, "Proyección mensual (12 meses desde la reapertura)", "Comisiones solo sobre eventos vía planner. Pauta como costo fijo mensual.")
cols = ["Mes", "N.º", "Temporada", "Ocup. V-S-D", "Ocup. L-J", "Eventos V-S-D", "Eventos L-J", "Eventos", "Ingreso sin IGV", "IGV (informativo)", "Comisiones", "Costos directos", "Pauta", "Contribución", "OPEX fijo", "EBITDA", "Acumulado vs CAPEX"]
header(mo, 4, cols, [10, 6, 11, 11, 10, 12, 11, 9, 15, 14, 13, 14, 11, 15, 12, 14, 18])
months = [(2026, 12)] + [(2027, m) for m in range(1, 12)]
MES = ["ene", "feb", "mar", "abr", "may", "jun", "jul", "ago", "set", "oct", "nov", "dic"]
for i, (y, m) in enumerate(months):
    r = 5 + i
    mo.cell(row=r, column=1, value=f"{MES[m - 1]}-{str(y)[2:]}")
    mo.cell(row=r, column=2, value=i + 1)
    high = m in (12, 1, 2, 3)
    mo.cell(row=r, column=3, value="Alta" if high else "Baja")
    mo.cell(row=r, column=4, value=f"=IF(B{r}<=3,{P['o_vsd3']},{P['o_vsd6']})").number_format = PCT
    mo.cell(row=r, column=5, value=f"=ROUND({P['o_lj6']}*MIN(1,B{r}/6),3)").number_format = PCT
    mo.cell(row=r, column=6, value=f"={P['d_vsd']}*D{r}").number_format = NUM
    mo.cell(row=r, column=7, value=f"=({P['d_lj']}-{P['d_block']})*E{r}").number_format = NUM
    mo.cell(row=r, column=8, value=f"=F{r}+G{r}").number_format = NUM
    factor = f"IF(C{r}=\"Alta\",{P['high']},1)"
    mo.cell(row=r, column=9, value=f"=(F{r}*{SS['t_vsd']}+G{r}*{SS['t_lj']})*{factor}").number_format = PEN
    mo.cell(row=r, column=10, value=f"=I{r}*{P['igv']}").number_format = PEN
    mo.cell(row=r, column=11, value=f"=I{r}*{P['ch_pl']}*{P['com']}").number_format = PEN
    mo.cell(row=r, column=12, value=f"=H{r}*{P['dc']}").number_format = PEN
    mo.cell(row=r, column=13, value=f"={P['ads']}").number_format = PEN
    mo.cell(row=r, column=14, value=f"=I{r}-K{r}-L{r}-M{r}").number_format = PEN
    mo.cell(row=r, column=15, value=f"={OPEX_SUM}").number_format = PEN
    mo.cell(row=r, column=16, value=f"=N{r}-O{r}").number_format = PEN
    prev = f"Q{r - 1}" if i else f"-({CAPEX_SUM})"
    mo.cell(row=r, column=17, value=f"={prev}+P{r}").number_format = PEN
tr = 5 + len(months)
mo.cell(row=tr, column=1, value="Total año").font = Font(bold=True)
for col in range(6, 17):
    L = get_column_letter(col)
    c = mo.cell(row=tr, column=col, value=f"=SUM({L}5:{L}{tr - 1})")
    c.font = Font(bold=True)
    c.number_format = NUM if col <= 8 else PEN
mo.cell(row=tr + 2, column=1, value="Mes de payback (dentro del año)").font = Font(bold=True)
mo.cell(row=tr + 2, column=6, value=f'=IFERROR(INDEX(A5:A{tr - 1},MATCH(TRUE,INDEX(Q5:Q{tr - 1}>=0,0),0)),"Después del mes 12")')
mo.freeze_panes = "C5"

# ---------------------------------------------------------------------------
# Equilibrio
# ---------------------------------------------------------------------------
eq = wb.create_sheet("Equilibrio")
title(eq, "Punto de equilibrio y payback (estado estable)")
header(eq, 4, ["Indicador", "Valor", "Cómo se calcula"], [48, 16, 70])
EQ = [
    ("Ticket promedio ponderado (sin IGV)", f"={SS['ticket']}", PEN, "Unit economics"),
    ("Contribución por evento (promedio de canales, sin pauta)", f"=B5-{P['dc']}-B5*{P['ch_pl']}*{P['com']}", PEN, "Ticket − costo directo − comisión media"),
    ("Costos fijos mensuales (OPEX + pauta)", f"={OPEX_SUM}+{P['ads']}", PEN, "Parámetros"),
    ("Eventos de equilibrio por mes", "=IF(B6>0,B7/B6,0)", NUM, "Costos fijos / contribución por evento"),
    ("Días disponibles por mes", f"={P['d_vsd']}+{P['d_lj']}-{P['d_block']}", NUM, "V-S-D + L-J − mantenimiento"),
    ("Ocupación de equilibrio", "=IF(B9>0,B8/B9,0)", PCT, "Eventos de equilibrio / días disponibles"),
    ("Eventos por mes en meta (estado estable)", f"={SS['e_tot']}", NUM, "Metas S10"),
    ("EBITDA mensual en meta", "=B11*B6-B7", PEN, "Eventos × contribución − fijos"),
    ("CAPEX total", f"={CAPEX_SUM}", PEN, "Parámetros"),
    ("Payback del CAPEX (meses en estado estable)", '=IF(B12>0,B13/B12,"No recupera")', NUM, "CAPEX / EBITDA mensual"),
    ("Ingreso neto por día disponible en meta", "=IF(B9>0,(B11*B5-B11*B5*" + P["ch_pl"] + "*" + P["com"] + "-B11*" + P["dc"] + ")/B9,0)", PEN, "KPI del dashboard"),
]
for i, (label, formula, fmt, how) in enumerate(EQ):
    r = 5 + i
    eq.cell(row=r, column=1, value=label)
    c = eq.cell(row=r, column=2, value=formula)
    c.number_format = fmt
    c.font = Font(bold=True)
    eq.cell(row=r, column=3, value=how)

# ---------------------------------------------------------------------------
# Sensibilidad
# ---------------------------------------------------------------------------
se = wb.create_sheet("Sensibilidad")
title(se, "Sensibilidad: EBITDA mensual (S/) según ocupación total y ticket", "Ocupación sobre todos los días disponibles; ticket = promedio ponderado × (1 + variación).")
se.cell(row=4, column=1, value="Ocupación ↓ / Ticket →").font = Font(bold=True)
deltas = [-0.2, -0.1, 0, 0.1, 0.2]
occs = [0.2, 0.3, 0.4, 0.5, 0.6, 0.7, 0.8]
for j, d in enumerate(deltas):
    c = se.cell(row=4, column=2 + j, value=d)
    c.number_format = '+0%;-0%;0%'
    c.fill = HEAD
    c.font = Font(bold=True, color=SAND)
for i, o in enumerate(occs):
    r = 5 + i
    c = se.cell(row=r, column=1, value=o)
    c.number_format = "0%"
    c.fill = HEAD
    c.font = Font(bold=True, color=SAND)
    for j in range(len(deltas)):
        L = get_column_letter(2 + j)
        t = f"({SS['ticket']}*(1+{L}$4))"
        f = f"=$A{r}*Equilibrio!$B$9*({t}-{P['dc']}-{t}*{P['ch_pl']}*{P['com']})-Equilibrio!$B$7"
        se.cell(row=r, column=2 + j, value=f).number_format = PEN
se.conditional_formatting.add(
    f"B5:{get_column_letter(1 + len(deltas))}{4 + len(occs)}",
    ColorScaleRule(start_type="min", start_color="E8A598", mid_type="num", mid_value=0, mid_color="FFFFFF", end_type="max", end_color="9CC5A1"),
)
for col in range(1, 7):
    se.column_dimensions[get_column_letter(col)].width = 18

# ---------------------------------------------------------------------------
# Pauta vs planner
# ---------------------------------------------------------------------------
pv = wb.create_sheet("Pauta vs planner")
title(pv, "Punto de equilibrio entre pauta y planner", "Si el CAC directo supera la comisión del planner, cada evento por pauta deja menos margen que uno por planner.")
header(pv, 4, ["Indicador", "Valor", "Nota"], [52, 16, 62])
PV = [
    ("Ticket promedio ponderado (sin IGV)", f"={SS['ticket']}", PEN, ""),
    ("Comisión del planner por evento (= CAC máximo tolerable)", f"=B5*{P['com']}", PEN, "Equilibrio: pagar más que esto por cliente directo no conviene"),
    ("CAC directo actual (supuesto)", f"={SS['cac']}", PEN, "Pauta mensual / eventos directos"),
    ("Veredicto", '=IF(B7<=B6,"La pauta cuesta menos que la comisión: OK","La pauta cuesta MÁS que la comisión: revisar campañas")', None, "Mismo criterio que la alerta del dashboard"),
    ("CPL máximo con la conversión supuesta", f"=B6*{P['conv']}", PEN2, "CAC máximo × conversión lead → evento"),
    ("Eventos directos necesarios para justificar la pauta", f"=IF(B6>0,{P['ads']}/B6,0)", NUM, "Pauta / CAC máximo"),
]
for i, (label, formula, fmt, note) in enumerate(PV):
    r = 5 + i
    pv.cell(row=r, column=1, value=label)
    c = pv.cell(row=r, column=2, value=formula)
    if fmt:
        c.number_format = fmt
    c.font = Font(bold=True)
    pv.cell(row=r, column=3, value=note)
r = 5 + len(PV) + 1
pv.cell(row=r, column=1, value="CPL máximo según la conversión lead → evento").font = Font(bold=True, color=GREEN)
r += 1
for v in (0.02, 0.04, 0.06, 0.08, 0.10, 0.15):
    c = pv.cell(row=r, column=1, value=v)
    c.number_format = "0%"
    pv.cell(row=r, column=2, value=f"=$B$6*A{r}").number_format = PEN2
    r += 1

OUT.parent.mkdir(parents=True, exist_ok=True)
wb.save(OUT)
print(f"✓ {OUT.relative_to(ROOT)}")
